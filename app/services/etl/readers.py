"""Tabular readers: CSV / TSV / TXT / XLSX, with encoding and delimiter detection.

Every reader yields plain lists of strings so the mapping layer never cares about the format.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterator

csv.field_size_limit(16 * 1024 * 1024)

SUPPORTED = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm"}


@dataclass
class Table:
    headers: list[str]
    meta: dict = field(default_factory=dict)
    _iter_factory: object = None
    _progress: object = None

    def rows(self) -> Iterator[tuple[int, list[str]]]:
        return self._iter_factory()

    def progress(self) -> float:
        return self._progress() if self._progress else 0.0


def _decode(raw: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace"), "latin-1"


def _delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    head = text[:65536]
    try:
        return csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: head.count(d) for d in ",;\t|"}
        return max(counts, key=counts.get)


def _headers(raw: list[str]) -> list[str]:
    out = []
    for i, h in enumerate(raw):
        h = (h or "").replace("﻿", "").strip()
        out.append(h if h else f"column_{i + 1}")
    return out


def _open_csv(path: Path) -> Table:
    text, encoding = _decode(path.read_bytes())
    delim = _delimiter(text, path.suffix.lower())
    headers = _headers(next(csv.reader(io.StringIO(text), delimiter=delim), []))
    total = max(len(text), 1)
    state = {"buffer": io.StringIO(text)}

    def iterate():  # each call is a fresh pass over the file
        buffer = state["buffer"] = io.StringIO(text)
        reader = csv.reader(buffer, delimiter=delim)
        next(reader, None)
        for row_no, row in enumerate(reader, start=2):
            if not any(cell.strip() for cell in row):
                continue
            yield row_no, row

    return Table(headers, {"format": "csv", "encoding": encoding, "delimiter": repr(delim)},
                 iterate, lambda: state["buffer"].tell() / total)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _open_xlsx(path: Path) -> Table:
    from openpyxl import load_workbook  # optional dependency, only needed for Excel files

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = max(wb.worksheets, key=lambda s: s.max_row or 0)
    it = ws.iter_rows(values_only=True)
    headers = _headers([_cell(v) for v in next(it, [])])
    total = max(ws.max_row or 1, 1)
    state = {"row": 1}

    def iterate():  # each call is a fresh pass over the sheet
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row_no, values in enumerate(rows, start=2):
            state["row"] = row_no
            row = [_cell(v) for v in values]
            if any(c.strip() for c in row):
                yield row_no, row

    return Table(headers, {"format": "xlsx", "sheet": ws.title}, iterate, lambda: state["row"] / total)


def open_table(path: str | Path) -> Table:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type {suffix}. Use one of: {', '.join(sorted(SUPPORTED))}.")
    return _open_xlsx(path) if suffix in {".xlsx", ".xlsm"} else _open_csv(path)


def sample(path: str | Path, limit: int = 40) -> tuple[list[str], list[list[str]], dict]:
    table = open_table(path)
    rows = []
    for _, row in table.rows():
        rows.append(row)
        if len(rows) >= limit:
            break
    width = max([len(table.headers)] + [len(r) for r in rows])
    headers = table.headers + [f"column_{i + 1}" for i in range(len(table.headers), width)]
    return headers, rows, table.meta
